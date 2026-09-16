import os
import json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    """Sets background color for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_docs():
    base_dir = r"d:\zoolearn\zoolearn(10.7.26)\zoolearn"
    all_animals_path = os.path.join(base_dir, "allAnimalData.json")
    all_phylum_path = os.path.join(base_dir, "AllPhylumData.json")
    
    # 1. Parse Species and Phylum data
    phylum_species_map = {}
    all_species_count = 0

    # Load from allAnimalData.json
    if os.path.exists(all_animals_path):
        with open(all_animals_path, 'r', encoding='utf-8') as f:
            animals_data = json.load(f)
            # could be list of phylums with animals
            if isinstance(animals_data, list):
                for p_item in animals_data:
                    p_name = p_item.get("phylum", "").lower()
                    if not p_name:
                        p_name = p_item.get("slug", "").lower()
                    if p_name:
                        if p_name not in phylum_species_map:
                            phylum_species_map[p_name] = {}
                        for sp in p_item.get("animals", []):
                            sp_slug = sp.get("slug", "").strip().lower().replace(" ", "-")
                            sp_name = sp.get("name", sp_slug.capitalize())
                            if sp_slug:
                                phylum_species_map[p_name][sp_slug] = sp_name

    # Load from AllPhylumData.json for completeness
    if os.path.exists(all_phylum_path):
        with open(all_phylum_path, 'r', encoding='utf-8') as f:
            phylum_data = json.load(f)
            if isinstance(phylum_data, dict):
                for key, item in phylum_data.items():
                    if isinstance(item, dict):
                        phylum = item.get("phylum_source", "").strip().lower()
                        slug = item.get("slug", "").strip().lower()
                        name = item.get("name", slug.capitalize())
                        if phylum and slug:
                            slug_clean = slug.replace(" ", "-")
                            if phylum not in phylum_species_map:
                                phylum_species_map[phylum] = {}
                            phylum_species_map[phylum][slug_clean] = name

    # Count total species
    for p, sp_dict in phylum_species_map.items():
        all_species_count += len(sp_dict)

    # 2. Build Document
    doc = Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Palette
    COLOR_PRIMARY = RGBColor(13, 143, 69)      # Emerald #0D8F45
    COLOR_DARK = RGBColor(15, 23, 42)          # Slate 900 #0F172A
    COLOR_MUTED = RGBColor(100, 116, 139)      # Slate 500 #64748B
    HEX_HEADER_BG = "0D8F45"
    HEX_LIGHT_ROW = "F8FAFC"
    HEX_ALT_ROW = "FFFFFF"

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("ZooLearn Platform — PRD Route Specifications")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # Subtitle / Meta
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(20)
    run_sub = sub_p.add_run("Complete Reference of Web Application Page Routes & REST API Endpoints\n"
                           f"Base Production URL: https://zoolearn.in  |  Local Development: http://localhost:3000\n"
                           f"Total Documented Phylums: {len(phylum_species_map)}  |  Total Documented Species: {all_species_count}")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(10.5)
    run_sub.font.color.rgb = COLOR_MUTED

    # Section 1: Core Web Routes
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    r_h1 = h1.add_run("1. Core Web Application Routes")
    r_h1.font.name = "Calibri"
    r_h1.font.size = Pt(16)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_DARK

    core_routes = [
        ("Home / Landing Page", "/", "https://zoolearn.in/", "Public", "Main landing experience with Hero, 3D animations, Conceptual Learning, and Modules."),
        ("ZooHub Catalog Root", "/zoohub", "https://zoolearn.in/zoohub", "Public", "Primary taxonomy portal with Phylum classification cards and instant search."),
        ("User Dashboard", "/dashboard", "https://zoolearn.in/dashboard", "Protected", "Personalized user analytics, saved taxonomy bookmarks, and learning track progress."),
        ("User Login / Auth", "/login", "https://zoolearn.in/login", "Public", "Authentication portal for student and educator access."),
        ("Conceptual Learning Hub", "/conceptual-learning", "https://zoolearn.in/conceptual-learning", "Public", "Interactive biological modules and concept deep-dives."),
        ("11th Standard Module", "/modules/11th", "https://zoolearn.in/modules/11th", "Public", "Foundational biology curriculum tailored for Grade 11 students."),
        ("12th Standard Module", "/modules/12th", "https://zoolearn.in/modules/12th", "Public", "Advanced genetics, human physiology, and Grade 12 board preparation."),
        ("NEET Aspirant Module", "/modules/neet", "https://zoolearn.in/modules/neet", "Public", "High-yield competitive exam prep, question banks, and mock exams.")
    ]

    t1 = doc.add_table(rows=1, cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.autofit = False

    col_widths_1 = [Inches(1.8), Inches(1.8), Inches(0.8), Inches(2.1)]
    headers_1 = ["Page / Section", "Full Production URL", "Access", "Description"]
    
    # Format header row
    hdr_cells = t1.rows[0].cells
    for idx, name in enumerate(headers_1):
        hdr_cells[idx].text = name
        hdr_cells[idx].width = col_widths_1[idx]
        set_cell_background(hdr_cells[idx], HEX_HEADER_BG)
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[idx].paragraphs[0]
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, (page, path, url, access, desc) in enumerate(core_routes):
        row_cells = t1.add_row().cells
        bg_color = HEX_LIGHT_ROW if r_idx % 2 == 0 else HEX_ALT_ROW
        for c_idx, val in enumerate([page, url, access, desc]):
            row_cells[c_idx].text = val
            row_cells[c_idx].width = col_widths_1[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.name = "Calibri"
                p.runs[0].font.size = Pt(9.5)
                if c_idx == 1:
                    p.runs[0].font.color.rgb = COLOR_PRIMARY
                    p.runs[0].font.size = Pt(9)
                elif c_idx == 0:
                    p.runs[0].font.bold = True

    # Section 2: Phylum Catalog Landing Routes
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(24)
    h2.paragraph_format.space_after = Pt(8)
    r_h2 = h2.add_run("2. Phylum Landing Page Routes")
    r_h2.font.name = "Calibri"
    r_h2.font.size = Pt(16)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_DARK

    t2 = doc.add_table(rows=1, cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False
    col_widths_2 = [Inches(1.8), Inches(2.2), Inches(1.0), Inches(1.5)]
    headers_2 = ["Phylum Name", "Production URL", "Total Species", "Pattern"]

    hdr_cells_2 = t2.rows[0].cells
    for idx, name in enumerate(headers_2):
        hdr_cells_2[idx].text = name
        hdr_cells_2[idx].width = col_widths_2[idx]
        set_cell_background(hdr_cells_2[idx], HEX_HEADER_BG)
        set_cell_margins(hdr_cells_2[idx], top=120, bottom=120, left=120, right=120)
        p = hdr_cells_2[idx].paragraphs[0]
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    sorted_phylums = sorted(phylum_species_map.keys())
    for r_idx, phylum in enumerate(sorted_phylums):
        sp_count = len(phylum_species_map[phylum])
        p_url = f"https://zoolearn.in/zoohub/{phylum}"
        row_cells = t2.add_row().cells
        bg_color = HEX_LIGHT_ROW if r_idx % 2 == 0 else HEX_ALT_ROW
        for c_idx, val in enumerate([phylum.capitalize(), p_url, str(sp_count), f"/zoohub/{phylum}"]):
            row_cells[c_idx].text = val
            row_cells[c_idx].width = col_widths_2[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.name = "Calibri"
                p.runs[0].font.size = Pt(9.5)
                if c_idx == 1:
                    p.runs[0].font.color.rgb = COLOR_PRIMARY
                elif c_idx == 0:
                    p.runs[0].font.bold = True

    # Section 3: ZooHub Detailed Species URLs
    doc.add_page_break()
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(8)
    r_h3 = h3.add_run("3. Complete ZooHub Species Detail Routes")
    r_h3.font.name = "Calibri"
    r_h3.font.size = Pt(16)
    r_h3.font.bold = True
    r_h3.font.color.rgb = COLOR_DARK

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(14)
    r_p_intro = p_intro.add_run(
        "The following section provides the exact production URL routes for each individual species page, "
        "structured by Phylum hierarchy (Pattern: https://zoolearn.in/zoohub/[phylum]/[species])."
    )
    r_p_intro.font.name = "Calibri"
    r_p_intro.font.size = Pt(10)
    r_p_intro.font.color.rgb = COLOR_MUTED

    for phylum in sorted_phylums:
        phylum_title_p = doc.add_paragraph()
        phylum_title_p.paragraph_format.space_before = Pt(14)
        phylum_title_p.paragraph_format.space_after = Pt(6)
        r_pt = phylum_title_p.add_run(f"Phylum: {phylum.capitalize()} ({len(phylum_species_map[phylum])} Species)")
        r_pt.font.name = "Calibri"
        r_pt.font.size = Pt(13)
        r_pt.font.bold = True
        r_pt.font.color.rgb = COLOR_PRIMARY

        t_sp = doc.add_table(rows=1, cols=3)
        t_sp.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_sp.autofit = False
        col_widths_sp = [Inches(1.8), Inches(1.5), Inches(3.2)]
        headers_sp = ["Species Name", "Slug", "Production Route URL"]

        hdr_cells_sp = t_sp.rows[0].cells
        for idx, name in enumerate(headers_sp):
            hdr_cells_sp[idx].text = name
            hdr_cells_sp[idx].width = col_widths_sp[idx]
            set_cell_background(hdr_cells_sp[idx], "1E293B") # Dark slate header
            set_cell_margins(hdr_cells_sp[idx], top=100, bottom=100, left=100, right=100)
            p = hdr_cells_sp[idx].paragraphs[0]
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

        species_dict = phylum_species_map[phylum]
        sorted_species_slugs = sorted(species_dict.keys())

        for s_idx, sp_slug in enumerate(sorted_species_slugs):
            sp_name = species_dict[sp_slug]
            sp_url = f"https://zoolearn.in/zoohub/{phylum}/{sp_slug}"
            row_cells = t_sp.add_row().cells
            bg_color = HEX_LIGHT_ROW if s_idx % 2 == 0 else HEX_ALT_ROW
            for c_idx, val in enumerate([sp_name, sp_slug, sp_url]):
                row_cells[c_idx].text = val
                row_cells[c_idx].width = col_widths_sp[c_idx]
                set_cell_background(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
                p = row_cells[c_idx].paragraphs[0]
                if len(p.runs) > 0:
                    p.runs[0].font.name = "Calibri"
                    p.runs[0].font.size = Pt(9)
                    if c_idx == 2:
                        p.runs[0].font.color.rgb = COLOR_PRIMARY
                    elif c_idx == 0:
                        p.runs[0].font.bold = True

    # Section 4: Backend API Endpoints
    doc.add_page_break()
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(8)
    r_h4 = h4.add_run("4. Backend API Endpoints (Next.js Route Handlers)")
    r_h4.font.name = "Calibri"
    r_h4.font.size = Pt(16)
    r_h4.font.bold = True
    r_h4.font.color.rgb = COLOR_DARK

    api_routes = [
        ("POST", "/api/auth/login", "https://zoolearn.in/api/auth/login", "Authenticate user credentials & issue session token"),
        ("POST", "/api/auth/register", "https://zoolearn.in/api/auth/register", "Register new user / student account"),
        ("POST", "/api/auth/logout", "https://zoolearn.in/api/auth/logout", "Revoke user session & clear cookies"),
        ("GET", "/api/auth/me", "https://zoolearn.in/api/auth/me", "Fetch currently authenticated user profile")
    ]

    t_api = doc.add_table(rows=1, cols=4)
    t_api.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_api.autofit = False
    col_widths_api = [Inches(1.0), Inches(2.0), Inches(2.2), Inches(1.3)]
    headers_api = ["Method", "Endpoint Path", "Full URL", "Functionality"]

    hdr_cells_api = t_api.rows[0].cells
    for idx, name in enumerate(headers_api):
        hdr_cells_api[idx].text = name
        hdr_cells_api[idx].width = col_widths_api[idx]
        set_cell_background(hdr_cells_api[idx], HEX_HEADER_BG)
        set_cell_margins(hdr_cells_api[idx], top=120, bottom=120, left=120, right=120)
        p = hdr_cells_api[idx].paragraphs[0]
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for a_idx, (method, endpoint, url, desc) in enumerate(api_routes):
        row_cells = t_api.add_row().cells
        bg_color = HEX_LIGHT_ROW if a_idx % 2 == 0 else HEX_ALT_ROW
        for c_idx, val in enumerate([method, endpoint, url, desc]):
            row_cells[c_idx].text = val
            row_cells[c_idx].width = col_widths_api[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.name = "Calibri"
                p.runs[0].font.size = Pt(9.5)
                if c_idx == 0:
                    p.runs[0].font.bold = True
                    p.runs[0].font.color.rgb = COLOR_PRIMARY
                elif c_idx == 2:
                    p.runs[0].font.color.rgb = COLOR_PRIMARY
                    p.runs[0].font.size = Pt(8.5)

    # Save to documents
    docs_dir = os.path.join(base_dir, "docs")
    os.makedirs(docs_dir, exist_ok=True)
    
    out_docx_1 = os.path.join(docs_dir, "ZooLearn_API_and_Page_Routes_PRD.docx")
    out_docx_2 = os.path.join(base_dir, "ZooLearn_API_and_Page_Routes_PRD.docx")
    
    doc.save(out_docx_1)
    doc.save(out_docx_2)
    print(f"Successfully generated DOCX at: {out_docx_1} and {out_docx_2}")

    # Also generate comprehensive Markdown version
    md_path = os.path.join(docs_dir, "API_ROUTES_PRD.md")
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write("# ZooLearn Platform — PRD Route Specifications\n\n")
        f.write("> **Base Production Domain:** `https://zoolearn.in`  \n")
        f.write("> **Local Development Domain:** `http://localhost:3000`  \n")
        f.write(f"> **Total Documented Phylums:** {len(phylum_species_map)} | **Total Documented Species:** {all_species_count}\n\n")
        
        f.write("## 1. Core Web Application Routes\n\n")
        f.write("| Page / Section | Path | Production URL | Access | Description |\n")
        f.write("| :--- | :--- | :--- | :--- | :--- |\n")
        for page, path, url, access, desc in core_routes:
            f.write(f"| **{page}** | `{path}` | [{url}]({url}) | `{access}` | {desc} |\n")
        
        f.write("\n## 2. Phylum Landing Page Routes\n\n")
        f.write("| Phylum Name | Production URL | Total Species | Pattern |\n")
        f.write("| :--- | :--- | :--- | :--- |\n")
        for phylum in sorted_phylums:
            sp_count = len(phylum_species_map[phylum])
            p_url = f"https://zoolearn.in/zoohub/{phylum}"
            f.write(f"| **{phylum.capitalize()}** | [{p_url}]({p_url}) | {sp_count} | `/zoohub/{phylum}` |\n")
            
        f.write("\n## 3. Complete ZooHub Species Detail Routes\n\n")
        for phylum in sorted_phylums:
            species_dict = phylum_species_map[phylum]
            sorted_species_slugs = sorted(species_dict.keys())
            f.write(f"### Phylum: {phylum.capitalize()} ({len(sorted_species_slugs)} Species)\n\n")
            f.write("| Species Name | Slug | Full Production Route URL |\n")
            f.write("| :--- | :--- | :--- |\n")
            for sp_slug in sorted_species_slugs:
                sp_name = species_dict[sp_slug]
                sp_url = f"https://zoolearn.in/zoohub/{phylum}/{sp_slug}"
                f.write(f"| **{sp_name}** | `{sp_slug}` | [{sp_url}]({sp_url}) |\n")
            f.write("\n")
            
        f.write("## 4. Backend API Endpoints\n\n")
        f.write("| Method | Endpoint Path | Production URL | Functionality |\n")
        f.write("| :--- | :--- | :--- | :--- |\n")
        for method, endpoint, url, desc in api_routes:
            f.write(f"| `{method}` | `{endpoint}` | [{url}]({url}) | {desc} |\n")

    print(f"Successfully generated Markdown at: {md_path}")

if __name__ == "__main__":
    generate_docs()
